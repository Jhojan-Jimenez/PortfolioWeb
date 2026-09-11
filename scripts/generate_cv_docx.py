import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def add_bottom_border(paragraph, color_hex="94A3B8", size="6"):
    """Adds a clean bottom border to section headings (ATS-safe native XML)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{size}" w:space="3" w:color="{color_hex}"/></w:pBdr>')
    pPr.append(pBdr)

def build_cv_es(output_path):
    doc = docx.Document()
    
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Pt(28)
    section.bottom_margin = Pt(28)
    section.left_margin = Pt(36)
    section.right_margin = Pt(36)
    
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(9.2)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    
    PRIMARY = RGBColor(15, 23, 42)     # Slate 900
    SECONDARY = RGBColor(71, 85, 105)  # Slate 600
    
    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = PRIMARY
        add_bottom_border(p, color_hex="94A3B8", size="6")
        return p

    def add_bullet(text_runs, space_after=1.6):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.06
        
        bullet_run = p.add_run("•  ")
        bullet_run.font.size = Pt(8.0)
        bullet_run.font.color.rgb = SECONDARY
        
        for text, bold, italic in text_runs:
            r = p.add_run(text)
            r.bold = bold
            r.italic = italic
            r.font.size = Pt(9.0)
            r.font.color.rgb = PRIMARY
        return p

    # 1. HEADER
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1)
    run_name = p_name.add_run("JHOJAN JIMENEZ")
    run_name.bold = True
    run_name.font.size = Pt(15.5)
    run_name.font.color.rgb = PRIMARY

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Software Engineer | Backend & Cloud Architecture | Full Stack Capabilities")
    run_title.bold = True
    run_title.font.size = Pt(9.5)
    run_title.font.color.rgb = SECONDARY

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(3)
    run_contact = p_contact.add_run("jhojanjimene@gmail.com   •   +57 320 328 2014   •   linkedin.com/in/jhojan-jimenez-dev   •   dev.jhojan.cloud   •   Bogotá, Colombia")
    run_contact.font.size = Pt(8.6)
    run_contact.font.color.rgb = SECONDARY

    # 2. PERFIL PROFESIONAL
    add_section_header("Perfil Profesional")
    p_summary = doc.add_paragraph()
    p_summary.paragraph_format.space_before = Pt(1)
    p_summary.paragraph_format.space_after = Pt(2.5)
    p_summary.paragraph_format.line_spacing = 1.08
    p_summary.add_run(
        "Ingeniero de Software enfocado en desarrollo Backend, Arquitectura Cloud y Sistemas de Datos con más de 2 años de experiencia en producción. "
        "Especializado en Python (FastAPI, Django), Node.js/TypeScript, Java y desarrollo Full Stack complementario (React, Next.js). "
        "Experiencia integrando pipelines de ingesta y analítica de datos operativos (SQL, Pandas, ETL), despliegues contenerizados en GCP/AWS con Docker y CI/CD, "
        "y sistemas multiagente con OpenAI y búsqueda vectorial. Primer lugar en Sabana Hack 2025."
    ).font.size = Pt(8.9)

    # 3. EXPERIENCIA LABORAL
    add_section_header("Experiencia Laboral")

    # Role 1: Blurealty
    p_role1 = doc.add_paragraph()
    p_role1.paragraph_format.space_before = Pt(2.5)
    p_role1.paragraph_format.space_after = Pt(1)
    p_role1.paragraph_format.keep_with_next = True
    r = p_role1.add_run("Desarrollador Backend & Cloud")
    r.bold = True
    r.font.size = Pt(9.3)
    r2 = p_role1.add_run(" — Blurealty")
    r2.bold = True
    r2.font.size = Pt(9.3)
    r3 = p_role1.add_run("   |   Julio 2025 – Actualidad   |   Remoto")
    r3.italic = True
    r3.font.size = Pt(8.8)
    r3.font.color.rgb = SECONDARY

    add_bullet([
        ("Diseñé e implementé ", True, False),
        ("la arquitectura cloud en GCP (Cloud Run, Cloud SQL) para sistema multiagente con OpenAI, integrando Vector Storage y similitudes cosenoidales para recuperación semántica de documentos en flujos de Q&A y captura de leads.", False, False)
    ])
    add_bullet([
        ("Estructuré ", True, False),
        ("pipelines de CI/CD automatizado con GitHub Actions y Docker multi-stage builds, eliminando discrepancias de entorno y garantizando despliegues continuos sin tiempo de inactividad (zero-downtime).", False, False)
    ])
    add_bullet([
        ("Modelé ", True, False),
        ("esquemas relacionales en PostgreSQL con Alembic, implementando optimización de consultas e indexación para ingesta y análisis de datos operativos con alta concurrencia.", False, False)
    ])
    add_bullet([
        ("Lideré ", True, False),
        ("el desarrollo frontend full stack en React/TypeScript elevando el puntaje de accesibilidad de 3.4 a 9.2 (auditoría WAVE / ADA) e instrumenté telemetría de eventos para analítica de conversión.", False, False)
    ])

    # Role 2: GovLab
    p_role2 = doc.add_paragraph()
    p_role2.paragraph_format.space_before = Pt(3)
    p_role2.paragraph_format.space_after = Pt(1)
    p_role2.paragraph_format.keep_with_next = True
    r = p_role2.add_run("Desarrollador Backend & Sistemas de Datos")
    r.bold = True
    r.font.size = Pt(9.3)
    r2 = p_role2.add_run(" — GovLab")
    r2.bold = True
    r2.font.size = Pt(9.3)
    r3 = p_role2.add_run("   |   Febrero 2025 – Julio 2025   |   Remoto")
    r3.italic = True
    r3.font.size = Pt(8.8)
    r3.font.color.rgb = SECONDARY

    add_bullet([
        ("Construí ", True, False),
        ("la plataforma centralizada de analítica e ingesta de datos agrícolas para Corpohass con FastAPI y PostgreSQL, consolidando datasets de producción, hectáreas y siembra para toma de decisiones y reportes estadísticos.", False, False)
    ])
    add_bullet([
        ("Diseñé ", True, False),
        ("APIs RESTful y flujos de transformación de datos (ETL) aplicando validación estricta con Pydantic, SQLAlchemy y mitigación de vulnerabilidades OWASP con migraciones Alembic.", False, False)
    ])
    add_bullet([
        ("Implementé ", True, False),
        ("pipelines asíncronos de ingesta documental en AWS S3 con URLs pre-firmadas y despachadores de notificaciones transaccionales vía Twilio y SendGrid.", False, False)
    ])

    # Role 3: Centro de Innovación UCTS
    p_role3 = doc.add_paragraph()
    p_role3.paragraph_format.space_before = Pt(3)
    p_role3.paragraph_format.space_after = Pt(1)
    p_role3.paragraph_format.keep_with_next = True
    r = p_role3.add_run("Desarrollador de Software & Datos Clínicos")
    r.bold = True
    r.font.size = Pt(9.3)
    r2 = p_role3.add_run(" — Centro de Innovación UCTS")
    r2.bold = True
    r2.font.size = Pt(9.3)
    r3 = p_role3.add_run("   |   Enero 2025 – Junio 2025   |   Presencial")
    r3.italic = True
    r3.font.size = Pt(8.8)
    r3.font.color.rgb = SECONDARY

    add_bullet([
        ("Desarrollé ", True, False),
        ("un motor clínico determinista de recomendación en Node.js/TypeScript, creando algoritmos de evaluación matricial sobre datasets del PAI (edades, contraindicaciones y patologías crónicas).", False, False)
    ])
    add_bullet([
        ("Traduje ", True, False),
        ("directrices normativas de salud pública y epidemiología en modelos lógicos de datos y reglas algorítmicas de cero ambigüedad, optimizando además el sitio web institucional.", False, False)
    ])

    # Role 4: TurboCupones
    p_role4 = doc.add_paragraph()
    p_role4.paragraph_format.space_before = Pt(3)
    p_role4.paragraph_format.space_after = Pt(1)
    p_role4.paragraph_format.keep_with_next = True
    r = p_role4.add_run("Desarrollador Web & Backend (Full Stack)")
    r.bold = True
    r.font.size = Pt(9.3)
    r2 = p_role4.add_run(" — TurboCupones")
    r2.bold = True
    r2.font.size = Pt(9.3)
    r3 = p_role4.add_run("   |   Junio 2024 – Octubre 2024   |   Remoto")
    r3.italic = True
    r3.font.size = Pt(8.8)
    r3.font.color.rgb = SECONDARY

    add_bullet([
        ("Desarrollé ", True, False),
        ("APIs REST con Django REST Framework para plataforma de cupones de alto tráfico, implementando autenticación JWT y optimización de consultas SQL con paginación por cursor (<100ms de latencia).", False, False)
    ])
    add_bullet([
        ("Colaboré ", True, False),
        ("en equipo ágil Scrum construyendo interfaces interactivas con Next.js y consumiendo microservicios backend con flujos de integración continua.", False, False)
    ])

    # 4. LOGROS DESTACADOS
    add_section_header("Logros Destacados")
    add_bullet([
        ("1er Lugar — Sabana Hack 2025 (Universidad de La Sabana): ", True, False),
        ("Diseñé e implementé en equipo, en 24 horas, un sistema de alertas en tiempo real con IA para la Cruz Roja Colombiana, procesando variables climáticas para mitigación de emergencias.", False, False)
    ], space_after=1)

    # 5. HABILIDADES TÉCNICAS
    add_section_header("Habilidades Técnicas")
    
    def add_skill_line(category, items):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.2)
        p.paragraph_format.line_spacing = 1.05
        bullet_run = p.add_run("•  ")
        bullet_run.font.size = Pt(8.0)
        bullet_run.font.color.rgb = SECONDARY
        r_cat = p.add_run(category + ": ")
        r_cat.bold = True
        r_cat.font.size = Pt(8.9)
        r_cat.font.color.rgb = PRIMARY
        r_items = p.add_run(items)
        r_items.font.size = Pt(8.9)
        r_items.font.color.rgb = PRIMARY

    add_skill_line("Backend", "Python (FastAPI, Django REST Framework), Node.js, TypeScript, Java, NestJS, Express.js")
    add_skill_line("Frontend", "React, Next.js, TypeScript, Tailwind CSS, HTML5, CSS3, consumo de APIs RESTful")
    add_skill_line("Bases de Datos & Analítica", "PostgreSQL, SQL, Pandas, Modelado de Datos, Pipelines ETL, Vector Storage (Embeddings, Similitud Cosenoidal)")
    add_skill_line("Cloud & DevOps", "GCP (Cloud Run, Cloud SQL), AWS S3, Docker, CI/CD (GitHub Actions), Railway")
    add_skill_line("Herramientas & Entornos", "Linux, Git, Postman, JWT, Pydantic, Alembic, Scrum, APIs RESTful")

    # 6. EDUCACIÓN & IDIOMAS
    add_section_header("Educación & Idiomas")
    
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_before = Pt(2)
    p_edu.paragraph_format.space_after = Pt(1)
    p_edu.paragraph_format.keep_with_next = True
    r_deg = p_edu.add_run("Grado en Ingeniería Informática")
    r_deg.bold = True
    r_deg.font.size = Pt(9.3)
    r_inst = p_edu.add_run(" — Universidad de La Sabana")
    r_inst.bold = True
    r_inst.font.size = Pt(9.3)
    r_date = p_edu.add_run("   |   2022 – 2026   |   Chía, Colombia")
    r_date.italic = True
    r_date.font.size = Pt(8.8)
    r_date.font.color.rgb = SECONDARY

    add_bullet([
        ("Beca & Rendimiento: ", True, False),
        ("Beneficiario de ", False, False),
        ("Beca de Excelencia Académica del 80%", True, False),
        (" · Promedio acumulado: ", False, False),
        ("4.4 / 5.0", True, False),
        (" (Énfasis en Arquitectura de Software y Sistemas Distribuidos).", False, False)
    ], space_after=1)

    add_bullet([
        ("Idiomas: ", True, False),
        ("Español (Nativo)   •   Inglés (B2 — Competencia profesional técnica y comunicación fluida).", False, False)
    ], space_after=0)

    doc.save(output_path)
    print(f"Saved CV to {output_path}")

def build_cv_en(output_path):
    doc = docx.Document()
    
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Pt(28)
    section.bottom_margin = Pt(28)
    section.left_margin = Pt(36)
    section.right_margin = Pt(36)
    
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(9.2)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    
    PRIMARY = RGBColor(15, 23, 42)
    SECONDARY = RGBColor(71, 85, 105)
    
    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = PRIMARY
        add_bottom_border(p, color_hex="94A3B8", size="6")
        return p

    def add_bullet(text_runs, space_after=1.6):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.06
        
        bullet_run = p.add_run("•  ")
        bullet_run.font.size = Pt(8.0)
        bullet_run.font.color.rgb = SECONDARY
        
        for text, bold, italic in text_runs:
            r = p.add_run(text)
            r.bold = bold
            r.italic = italic
            r.font.size = Pt(9.0)
            r.font.color.rgb = PRIMARY
        return p

    # 1. HEADER
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1)
    run_name = p_name.add_run("JHOJAN JIMENEZ")
    run_name.bold = True
    run_name.font.size = Pt(15.5)
    run_name.font.color.rgb = PRIMARY

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Software Engineer | Backend & Cloud Architecture | Full Stack Capabilities")
    run_title.bold = True
    run_title.font.size = Pt(9.5)
    run_title.font.color.rgb = SECONDARY

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(3)
    run_contact = p_contact.add_run("jhojanjimene@gmail.com   •   +57 320 328 2014   •   linkedin.com/in/jhojan-jimenez-dev   •   dev.jhojan.cloud   •   Bogotá, Colombia")
    run_contact.font.size = Pt(8.6)
    run_contact.font.color.rgb = SECONDARY

    # 2. SUMMARY
    add_section_header("Professional Summary")
    p_summary = doc.add_paragraph()
    p_summary.paragraph_format.space_before = Pt(1)
    p_summary.paragraph_format.space_after = Pt(2.5)
    p_summary.paragraph_format.line_spacing = 1.08
    p_summary.add_run(
        "Software Engineer focused on Backend Development, Cloud Architecture, and Data Systems with 2+ years of production experience. "
        "Proficient in Python (FastAPI, Django), Node.js/TypeScript, Java, and complementary Full Stack development (React, Next.js). "
        "Hands-on experience engineering data ingestion & operational analytics pipelines (SQL, Pandas, ETL), containerized GCP/AWS deployments with Docker and CI/CD, "
        "and multi-agent AI systems with OpenAI and vector search. 1st place at Sabana Hack 2025."
    ).font.size = Pt(8.9)

    # 3. EXPERIENCE
    add_section_header("Work Experience")

    # Role 1: Blurealty
    p_role1 = doc.add_paragraph()
    p_role1.paragraph_format.space_before = Pt(2.5)
    p_role1.paragraph_format.space_after = Pt(1)
    p_role1.paragraph_format.keep_with_next = True
    r = p_role1.add_run("Backend & Cloud Infrastructure Engineer")
    r.bold = True
    r.font.size = Pt(9.3)
    r2 = p_role1.add_run(" — Blurealty")
    r2.bold = True
    r2.font.size = Pt(9.3)
    r3 = p_role1.add_run("   |   July 2025 – Present   |   Remote")
    r3.italic = True
    r3.font.size = Pt(8.8)
    r3.font.color.rgb = SECONDARY

    add_bullet([
        ("Architected and deployed ", True, False),
        ("cloud infrastructure on GCP (Cloud Run, Cloud SQL) for an OpenAI multi-agent system, integrating Vector Storage and cosine similarity for semantic document retrieval in Q&A and lead-capture pipelines.", False, False)
    ])
    add_bullet([
        ("Structured ", True, False),
        ("automated CI/CD pipelines with GitHub Actions and Docker multi-stage builds, eliminating environment drift and ensuring zero-downtime continuous delivery.", False, False)
    ])
    add_bullet([
        ("Modeled ", True, False),
        ("relational PostgreSQL schemas with Alembic, optimizing SQL queries and indexing to handle high-concurrency ingestion and operational data analysis.", False, False)
    ])
    add_bullet([
        ("Led ", True, False),
        ("full-stack frontend refactoring in React/TypeScript, elevating web accessibility audit score from 3.4 to 9.2 (WAVE / ADA) and instrumenting telemetry events for conversion analytics.", False, False)
    ])

    # Role 2: GovLab
    p_role2 = doc.add_paragraph()
    p_role2.paragraph_format.space_before = Pt(3)
    p_role2.paragraph_format.space_after = Pt(1)
    p_role2.paragraph_format.keep_with_next = True
    r = p_role2.add_run("Backend & Data Systems Engineer")
    r.bold = True
    r.font.size = Pt(9.3)
    r2 = p_role2.add_run(" — GovLab")
    r2.bold = True
    r2.font.size = Pt(9.3)
    r3 = p_role2.add_run("   |   February 2025 – July 2025   |   Remote")
    r3.italic = True
    r3.font.size = Pt(8.8)
    r3.font.color.rgb = SECONDARY

    add_bullet([
        ("Built ", True, False),
        ("the centralized agricultural data analytics and ingestion platform for Corpohass with FastAPI and PostgreSQL, consolidating production, acreage, and harvest datasets for operational decision-making.", False, False)
    ])
    add_bullet([
        ("Engineered ", True, False),
        ("production RESTful APIs and ETL data transformations applying strict Pydantic validation, SQLAlchemy, and OWASP vulnerability mitigation with Alembic migrations.", False, False)
    ])
    add_bullet([
        ("Implemented ", True, False),
        ("asynchronous document ingestion pipelines on AWS S3 via pre-signed URLs and automated transactional notification dispatchers via Twilio and SendGrid.", False, False)
    ])

    # Role 3: UCTS
    p_role3 = doc.add_paragraph()
    p_role3.paragraph_format.space_before = Pt(3)
    p_role3.paragraph_format.space_after = Pt(1)
    p_role3.paragraph_format.keep_with_next = True
    r = p_role3.add_run("Software & Health Data Engineer")
    r.bold = True
    r.font.size = Pt(9.3)
    r2 = p_role3.add_run(" — UCTS Innovation Center")
    r2.bold = True
    r2.font.size = Pt(9.3)
    r3 = p_role3.add_run("   |   January 2025 – June 2025   |   On-site")
    r3.italic = True
    r3.font.size = Pt(8.8)
    r3.font.color.rgb = SECONDARY

    add_bullet([
        ("Engineered ", True, False),
        ("a deterministic clinical recommendation engine in Node.js/TypeScript, developing multidimensional matrix evaluation algorithms over complex PAI datasets (ages, contraindications, pathologies).", False, False)
    ])
    add_bullet([
        ("Translated ", True, False),
        ("public health regulations alongside epidemiologists into zero-ambiguity logical data models, while modernizing institutional web architecture.", False, False)
    ])

    # Role 4: TurboCupones
    p_role4 = doc.add_paragraph()
    p_role4.paragraph_format.space_before = Pt(3)
    p_role4.paragraph_format.space_after = Pt(1)
    p_role4.paragraph_format.keep_with_next = True
    r = p_role4.add_run("Web & Backend Engineer (Full Stack)")
    r.bold = True
    r.font.size = Pt(9.3)
    r2 = p_role4.add_run(" — TurboCupones")
    r2.bold = True
    r2.font.size = Pt(9.3)
    r3 = p_role4.add_run("   |   June 2024 – October 2024   |   Remote")
    r3.italic = True
    r3.font.size = Pt(8.8)
    r3.font.color.rgb = SECONDARY

    add_bullet([
        ("Developed ", True, False),
        ("high-throughput REST APIs with Django REST Framework for a coupon platform, implementing JWT authentication, SQL query tuning, and cursor-based pagination (<100ms response times).", False, False)
    ])
    add_bullet([
        ("Collaborated ", True, False),
        ("in an agile Scrum team delivering responsive Next.js interfaces consuming backend microservices under continuous integration workflows.", False, False)
    ])

    # 4. ACHIEVEMENTS
    add_section_header("Key Achievements")
    add_bullet([
        ("1st Place — Sabana Hack 2025 (Universidad de La Sabana): ", True, False),
        ("Architected and implemented in 24 hours an AI-powered real-time emergency alert system for the Colombian Red Cross, processing environmental data for disaster risk mitigation.", False, False)
    ], space_after=1)

    # 5. TECHNICAL SKILLS
    add_section_header("Technical Skills")
    
    def add_skill_line(category, items):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.2)
        p.paragraph_format.line_spacing = 1.05
        bullet_run = p.add_run("•  ")
        bullet_run.font.size = Pt(8.0)
        bullet_run.font.color.rgb = SECONDARY
        r_cat = p.add_run(category + ": ")
        r_cat.bold = True
        r_cat.font.size = Pt(8.9)
        r_cat.font.color.rgb = PRIMARY
        r_items = p.add_run(items)
        r_items.font.size = Pt(8.9)
        r_items.font.color.rgb = PRIMARY

    add_skill_line("Backend", "Python (FastAPI, Django REST Framework), Node.js, TypeScript, Java, NestJS, Express.js")
    add_skill_line("Frontend", "React, Next.js, TypeScript, Tailwind CSS, HTML5, CSS3, RESTful API consumption")
    add_skill_line("Databases & Analytics", "PostgreSQL, SQL, Pandas, Data Modeling, ETL Pipelines, Vector Storage (Embeddings, Cosine Similarity)")
    add_skill_line("Cloud & DevOps", "GCP (Cloud Run, Cloud SQL), AWS S3, Docker, CI/CD (GitHub Actions), Railway")
    add_skill_line("Tools & Environments", "Linux, Git, Postman, JWT, Pydantic, Alembic, Scrum, RESTful APIs")

    # 6. EDUCATION & LANGUAGES
    add_section_header("Education & Languages")
    
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_before = Pt(2)
    p_edu.paragraph_format.space_after = Pt(1)
    p_edu.paragraph_format.keep_with_next = True
    r_deg = p_edu.add_run("B.S. in Computer Engineering")
    r_deg.bold = True
    r_deg.font.size = Pt(9.3)
    r_inst = p_edu.add_run(" — Universidad de La Sabana")
    r_inst.bold = True
    r_inst.font.size = Pt(9.3)
    r_date = p_edu.add_run("   |   2022 – 2026   |   Chía, Colombia")
    r_date.italic = True
    r_date.font.size = Pt(8.8)
    r_date.font.color.rgb = SECONDARY

    add_bullet([
        ("Scholarship & Academic Merit: ", True, False),
        ("Recipient of ", False, False),
        ("80% Academic Excellence Scholarship", True, False),
        (" · Cumulative GPA: ", False, False),
        ("4.4 / 5.0", True, False),
        (" (Emphasis on Software Architecture and Distributed Systems).", False, False)
    ], space_after=1)

    add_bullet([
        ("Languages: ", True, False),
        ("Spanish (Native)   •   English (B2 — Professional working proficiency & technical documentation).", False, False)
    ], space_after=0)

    doc.save(output_path)
    print(f"Saved Resume to {output_path}")

if __name__ == '__main__':
    build_cv_es('public/resume/Jhojan_JimenezCV.docx')
    build_cv_en('public/resume/Jhojan_Jimenez_Resume.docx')
